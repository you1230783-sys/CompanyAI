"""隔離的單次分析入口；資料只走管道，輸出由 Rust 驗證後發布。

此程式不把 Python 的物件限制當成安全沙箱；真正邊界是 Windows
AppContainer（無網路能力）與 Job Object（單一程序、記憶體及逾時限制）。
"""
import contextlib
import ctypes
from ctypes import wintypes
import io
import json
import struct
import sys
import traceback


def require_container():
    """拒絕一般權限直接啟動，避免測試或錯誤路由解除隔離。"""
    token = wintypes.HANDLE()
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    advapi = ctypes.WinDLL("advapi32", use_last_error=True)
    # x64 HANDLE 不能沿用 ctypes 的預設 int 參數，否則偽控制代碼會溢位。
    kernel.GetCurrentProcess.restype = wintypes.HANDLE
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    advapi.OpenProcessToken.argtypes = [wintypes.HANDLE, wintypes.DWORD, ctypes.POINTER(wintypes.HANDLE)]
    advapi.GetTokenInformation.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, wintypes.DWORD, ctypes.POINTER(wintypes.DWORD)]
    if not advapi.OpenProcessToken(kernel.GetCurrentProcess(), 8, ctypes.byref(token)):
        raise RuntimeError("Cannot query process token")
    try:
        value = wintypes.DWORD()
        length = wintypes.DWORD()
        if not advapi.GetTokenInformation(token, 29, ctypes.byref(value), 4, ctypes.byref(length)) or value.value != 1:
            raise RuntimeError("Python worker requires AppContainer")
    finally:
        kernel.CloseHandle(token)


def frame(stream, payload):
    stream.write(struct.pack("<I", len(payload)))
    stream.write(payload)
    stream.flush()


class BoundedLog(io.StringIO):
    """只保留少量 print 輸出；大量資料應透過成果表交付。"""
    def write(self, text):
        remaining = max(0, 8000 - self.tell())
        super().write(text[:remaining])
        return len(text)


def analyze(request):
    """穩定資料介面：tables、texts、metadata、emit_table、emit_excel。"""
    import pandas as pd
    import numpy as np
    import openpyxl

    tables, texts, metadata = {}, {}, {}
    for item in request["inputs"]:
        name = item["name"]
        metadata[name] = {k: v for k, v in item.items() if k not in ("rows", "text", "hex")}
        if item["kind"] == "dataset":
            # 先用 object 保留文字、空白、布林與錯誤；模型依明確用途轉換欄位。
            tables[name] = pd.DataFrame(item["rows"], columns=item["columns"], dtype=object)
        elif item["kind"] == "csv":
            texts[name] = item["text"]
            tables[name] = pd.read_csv(io.StringIO(item["text"]), dtype=str, keep_default_na=False)
        elif item["kind"] == "xlsx":
            # 只接受 broker 核准的生成檔；公司原始 Office 必須先走 COM。
            book = openpyxl.load_workbook(io.BytesIO(bytes.fromhex(item["hex"])), read_only=True, data_only=True)
            try:
                metadata[name]["sheets"] = book.sheetnames
                for sheet in book:
                    values = sheet.iter_rows(values_only=True)
                    columns = next(values, ())
                    rows = []
                    for row in values:
                        if len(rows) >= 100000:
                            raise ValueError("XLSX exceeds 100000 rows; select a smaller source")
                        rows.append(row)
                    tables[name + "/" + sheet.title] = pd.DataFrame(rows, columns=columns, dtype=object)
            finally:
                book.close()
        else:
            texts[name] = item["text"]
    outputs = []

    def scalar(value):
        """JSON 不接受 NaN／Infinity，缺值保留 null，不當成零。"""
        if value is None or pd.isna(value):
            return None
        if isinstance(value, np.generic):
            value = value.item()
        if isinstance(value, (str, bool, int, float)):
            if isinstance(value, float) and not np.isfinite(value):
                raise ValueError("Infinite result is not supported")
            return value
        if hasattr(value, "isoformat"):
            return value.isoformat()
        raise ValueError("Unsupported cell type: " + type(value).__name__)

    def emit_table(name, data):
        """交付可供既有圖表工具讀取的 CSV 資料集，完整資料不回傳模型。"""
        if len(outputs) >= 8:
            raise ValueError("At most 8 outputs")
        df = pd.DataFrame(data)
        if len(df) > 100000 or not 1 <= len(df.columns) <= 16:
            raise ValueError("Table limit: 100000 rows / 16 columns")
        outputs.append({"kind": "table", "name": name, "columns": [str(c) for c in df.columns],
                        "rows": [[scalar(v) for v in row] for row in df.itertuples(index=False, name=None)]})

    def emit_excel(name, sheets):
        """以 openpyxl 產生多工作表；字串永遠寫成文字，不意外執行公式。"""
        if len(outputs) >= 8 or not 1 <= len(sheets) <= 12:
            raise ValueError("Output/workbook sheet limit exceeded")
        book = openpyxl.Workbook()
        book.remove(book.active)
        for title, data in sheets.items():
            df = pd.DataFrame(data)
            if len(df) > 100000 or not 1 <= len(df.columns) <= 32:
                raise ValueError("Excel limit: 100000 rows / 32 columns")
            sheet = book.create_sheet(str(title))
            for row in [[str(c) for c in df.columns]]:
                sheet.append(row)
            for row in df.itertuples(index=False, name=None):
                sheet.append([scalar(v) for v in row])
            for row in sheet:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = openpyxl.styles.Font(bold=True)
        buffer = io.BytesIO()
        book.save(buffer)
        if buffer.tell() > 8 * 1024 * 1024:
            raise ValueError("Excel output exceeds 8 MiB")
        # 同程序再開啟驗證容器與工作表，Rust 另核對實際寫入 bytes。
        check = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()), read_only=True)
        names = check.sheetnames
        check.close()
        outputs.append({"kind": "xlsx", "name": name, "hex": buffer.getvalue().hex(), "sheets": names})

    scope = {"pd": pd, "np": np, "openpyxl": openpyxl, "io": io,
             "tables": tables, "texts": texts, "metadata": metadata,
             "emit_table": emit_table, "emit_excel": emit_excel, "result": {}}
    log = BoundedLog()
    with contextlib.redirect_stdout(log), contextlib.redirect_stderr(log):
        exec(compile(request["code"], "<company-ai-analysis>", "exec"), scope)
    summary = json.dumps(scope.get("result", {}), ensure_ascii=False, allow_nan=False)
    if len(summary) > 12000:
        raise ValueError("result exceeds 12000 characters; use emit_table for full data")
    return {"ok": True, "summary": json.loads(summary), "stdout": log.getvalue(), "outputs": outputs,
            "versions": {"python": sys.version.split()[0], "pandas": pd.__version__,
                         "numpy": np.__version__, "openpyxl": openpyxl.__version__}}


def main():
    require_container()
    reader, writer = sys.stdin.buffer, sys.stdout.buffer
    frame(writer, b"appcontainer-ready")
    length = struct.unpack("<I", reader.read(4))[0]
    if length > 64 * 1024 * 1024:
        raise ValueError("Input exceeds 64 MiB")
    payload = reader.read(length)
    if len(payload) != length:
        raise ValueError("Incomplete input")
    try:
        response = analyze(json.loads(payload))
    except BaseException:
        response = {"ok": False, "error": traceback.format_exc(limit=6)[-8000:]}
    encoded = json.dumps(response, ensure_ascii=False, allow_nan=False).encode("utf-8")
    if len(encoded) > 32 * 1024 * 1024:
        encoded = b'{"ok":false,"error":"Python output exceeds 32 MiB"}'
    frame(writer, encoded)


if __name__ == "__main__":
    main()
